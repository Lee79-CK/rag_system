"""Office 파서 — docx/pptx/xlsx (docs/TDD.md §7.7 step 2).

- .docx : python-docx (문단 + 표 → Markdown)
- .pptx : python-pptx (슬라이드별 텍스트 프레임)
- .xlsx : openpyxl (시트별 → Markdown 표)
"""

from __future__ import annotations

from pathlib import Path

from docx import Document as DocxDocument
from openpyxl import load_workbook
from pptx import Presentation
from rag_schemas import ParsedDocument, ParsedElement, ParsedPage

from ingestion.parser.errors import ParseFailed


def _cells_to_markdown(rows: list[list[str]]) -> str:
    rows = [r for r in rows if any(cell.strip() for cell in r)]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    header = rows[0]
    lines = ["| " + " | ".join(header) + " |", "| " + " | ".join(["---"] * width) + " |"]
    for row in rows[1:]:
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)


def parse_docx(path: str | Path, *, document_id: str, version_id: str) -> ParsedDocument:
    p = Path(path)
    try:
        doc = DocxDocument(str(p))
    except Exception as exc:  # noqa: BLE001
        raise ParseFailed(f"DOCX 열기 실패(손상/암호화): {p.name}", code="CORRUPT") from exc

    elements: list[ParsedElement] = []
    char_count = 0
    for para in doc.paragraphs:
        txt = para.text.strip()
        if not txt:
            continue
        char_count += len(txt)
        style = (para.style.name or "").lower() if para.style else ""
        etype = "Title" if style.startswith("heading") or style == "title" else "NarrativeText"
        elements.append(ParsedElement(type=etype, text=txt, metadata={"style": style}))

    table_count = 0
    for table in doc.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        md = _cells_to_markdown(rows)
        if md:
            table_count += 1
            char_count += len(md)
            elements.append(ParsedElement(type="Table", text=md, metadata={}))

    if not elements:
        raise ParseFailed(f"빈 DOCX: {p.name}", code="EMPTY")

    page = ParsedPage(page=1, elements=elements)
    return ParsedDocument(
        document_id=document_id,
        version_id=version_id,
        parser="python-docx",
        pages=[page],
        stats={
            "page_count": 1,
            "char_count": char_count,
            "table_count": table_count,
            "image_count": 0,
        },
    )


def parse_pptx(path: str | Path, *, document_id: str, version_id: str) -> ParsedDocument:
    p = Path(path)
    try:
        prs = Presentation(str(p))
    except Exception as exc:  # noqa: BLE001
        raise ParseFailed(f"PPTX 열기 실패(손상/암호화): {p.name}", code="CORRUPT") from exc

    pages: list[ParsedPage] = []
    char_count = 0
    image_count = 0
    for idx, slide in enumerate(prs.slides, start=1):
        elements: list[ParsedElement] = []
        for shape in slide.shapes:
            if shape.shape_type == 13:  # PICTURE
                image_count += 1
            if not shape.has_text_frame:
                continue
            txt = "\n".join(
                run.text for para in shape.text_frame.paragraphs for run in para.runs
            ).strip()
            if not txt:
                txt = shape.text_frame.text.strip()
            if txt:
                char_count += len(txt)
                elements.append(
                    ParsedElement(type="NarrativeText", text=txt, metadata={"slide": idx})
                )
        pages.append(ParsedPage(page=idx, elements=elements))

    if char_count == 0:
        raise ParseFailed(f"빈 PPTX: {p.name}", code="EMPTY")

    return ParsedDocument(
        document_id=document_id,
        version_id=version_id,
        parser="python-pptx",
        pages=pages,
        stats={
            "page_count": len(pages),
            "char_count": char_count,
            "table_count": 0,
            "image_count": image_count,
        },
    )


def parse_xlsx(path: str | Path, *, document_id: str, version_id: str) -> ParsedDocument:
    p = Path(path)
    try:
        wb = load_workbook(filename=str(p), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ParseFailed(f"XLSX 열기 실패(손상/암호화): {p.name}", code="CORRUPT") from exc

    elements: list[ParsedElement] = []
    char_count = 0
    table_count = 0
    try:
        for ws in wb.worksheets:
            rows: list[list[str]] = []
            for row in ws.iter_rows(values_only=True):
                rows.append(["" if c is None else str(c).strip() for c in row])
            md = _cells_to_markdown(rows)
            if not md:
                continue
            table_count += 1
            char_count += len(md)
            # 시트명을 Title로 프리픽스하여 섹션 맥락 유지.
            elements.append(
                ParsedElement(type="Title", text=f"[시트] {ws.title}", metadata={"sheet": ws.title})
            )
            elements.append(ParsedElement(type="Table", text=md, metadata={"sheet": ws.title}))
    finally:
        wb.close()

    if not elements:
        raise ParseFailed(f"빈 XLSX: {p.name}", code="EMPTY")

    page = ParsedPage(page=1, elements=elements)
    return ParsedDocument(
        document_id=document_id,
        version_id=version_id,
        parser="openpyxl",
        pages=[page],
        stats={
            "page_count": 1,
            "char_count": char_count,
            "table_count": table_count,
            "image_count": 0,
        },
    )
